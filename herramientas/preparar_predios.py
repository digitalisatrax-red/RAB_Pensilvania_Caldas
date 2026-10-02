"""Predios adquiridos / por adquirir, ANONIMIZADOS (data/tablas/predios.json).

La exclusión de datos personales está en el código: solo se emiten las claves de la lista
blanca y, antes de escribir, se verifica que ningún valor sensible leído aparezca en la salida.
"""
import os
import re
import unicodedata
import openpyxl

LISTA_BLANCA = ("id", "nombre_predio", "estado", "area_sig_ha", "area_escritura_ha", "cobertura_ha")
# Lo que NUNCA debe salir (prueba 3 de la especificación)
LISTA_NEGRA = ("propietario", "propietario_anterior", "no. matricula", "matricula", "cod_catastral", "cedula_cat", "cedula")

CLASES = [("AVANZADO", "bosque_avanzado"), ("JOVEN", "bosque_joven"), ("PINO", "plantacion_pino"), ("PASTIZAL", "pastizal"),
          ("MATORRAL", "matorral"), ("SIN VEG", "sin_vegetacion"), ("VIAS", "vias"), ("DRENAJE", "drenajes"), ("CULTIVO", "cultivos")]


def _n(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.replace("\xa0", "").strip()
        if not v:
            return None
        try:
            return float(v.replace(",", "."))
        except ValueError:
            return None
    return float(v)


def _txt(v):
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip() if v is not None else ""


def _nombre(v):
    t = _txt(v)
    return t.title() if t.isupper() else t


def _coberturas(cab, fila):
    out = {}
    for i, c in enumerate(cab):
        if not isinstance(c, str):
            continue
        cu = unicodedata.normalize("NFKD", c).encode("ascii", "ignore").decode().upper()
        for clave, nombre in CLASES:
            if clave in cu and i + 1 < len(fila) and (cab[i + 1] or "") .strip() == "Ha":
                v = _n(fila[i + 1])
                if v:
                    out[nombre] = round(v, 2)
    return out


def preparar(origen):
    base = os.path.join(origen, "BD_PREDIOS", "BD_PREDIOS")
    predios, sensibles, aviso = [], set(), []
    # --- Adquiridos (hoja PREDIOS_RAB) ---
    wb = openpyxl.load_workbook(os.path.join(base, "PREDIOS_ADQUIRIDOS_COBERTURAS.xlsx"), data_only=True)
    filas = list(wb["PREDIOS_RAB"].iter_rows(values_only=True))
    cab = [c for c in filas[0]]
    ix = {(_txt(c) if c else ""): i for i, c in enumerate(cab)}
    omitidas = 0
    for f in filas[1:]:
        ident = _n(f[ix["ID_"]])
        if ident is None:
            omitidas += 1
            continue
        for campo in ("COD_CATASTRAL", "NO. MATRICULA", "PROPIETARIO_ANTERIOR", "PROPIETARIO"):
            v = _txt(f[ix[campo]])
            if len(v) >= 4:
                sensibles.add(v)
        predios.append(dict(id=f"ad-{int(ident)}", nombre_predio=_nombre(f[ix["NOM_PREDIO"]]), estado="Adquirido",
                            area_sig_ha=_n(f[ix["ÁREA_SIG_HA"]]), area_escritura_ha=_n(f[ix["ÁREA_ESCRITURA_Ha"]]),
                            cobertura_ha=_coberturas(cab, f)))
    if omitidas:
        aviso.append(f"{omitidas} fila(s) de PREDIOS_RAB sin ID numérico (subtotales) no se publican como predios.")
    # --- Por adquirir (Hoja1, primer bloque 'PREDIOS LISTOS PARA COMPRA') ---
    wb2 = openpyxl.load_workbook(os.path.join(base, "PREDIOS_ADQUIRIR.xlsx"), data_only=True)
    filas = list(wb2["Hoja1"].iter_rows(values_only=True))
    cab2 = None
    for i, f in enumerate(filas):
        if _txt(f[0]) == "ID":
            cab2 = f
            inicio = i + 1
            break
    for f in filas[inicio:]:
        ident = _n(f[0])
        if ident is None:
            break
        for k in (1, 3):
            v = _txt(f[k])
            if len(v) >= 4:
                sensibles.add(v)
        predios.append(dict(id=f"pa-{int(ident)}", nombre_predio=_nombre(f[2]), estado="Por adquirir (listo para compra)",
                            area_sig_ha=round(_n(f[4]), 2) if _n(f[4]) else None, area_escritura_ha=None,
                            cobertura_ha=_coberturas(cab2, f)))
    for p in predios:
        if p["area_sig_ha"] is not None:
            p["area_sig_ha"] = round(p["area_sig_ha"], 2)
        assert set(p) == set(LISTA_BLANCA), p
    ad = [p for p in predios if p["estado"] == "Adquirido"]
    pa = [p for p in predios if p["estado"] != "Adquirido"]
    resumen = dict(
        adquiridos=len(ad), por_adquirir=len(pa),
        ha_adquiridas_sig=round(sum(p["area_sig_ha"] or 0 for p in ad), 2),
        ha_por_adquirir_sig=round(sum(p["area_sig_ha"] or 0 for p in pa), 2),
        adquiridos_sin_area_sig=sum(1 for p in ad if p["area_sig_ha"] is None),
    )
    salida = dict(version=1, predios=predios, resumen=resumen,
                  nota="Datos anonimizados: solo nombre del predio, áreas y coberturas declaradas por FCV.")
    texto = str(salida)
    for s in sensibles:
        assert s not in texto, "dato personal en la salida de predios"
    return salida, aviso
